from pathlib import Path
import math
import json

import numpy as np
import pandas as pd
from scipy.stats import wilcoxon, norm
from openpyxl.styles import PatternFill, Font


# =============================================================================
# Paths
# =============================================================================

RESULTS_ROOT = Path(r"F:\Results\SAM_Benchmarking")

OUTPUT_DIR = (
    RESULTS_ROOT
    / "Model_comparison"
    / "statistical_comparison"
)

OUTPUT_XLSX = OUTPUT_DIR / "posneg_hybrid_vs_segformer_and_boxonly_statistics.xlsx"


# =============================================================================
# Methods
# =============================================================================

POSNEG_METHOD_FOLDER = "SegFormer_SurgiSAM2_AutoBox_PosNeg_Fallback"
POSNEG_TRAINING_STATE = "hybrid"
POSNEG_PROMPT_MODE = "Auto_box_posneg_fallback_dice_0p80"

BOXONLY_METHOD_FOLDER = "SegFormer_SurgiSAM2_AutoBox_Fallback"
BOXONLY_TRAINING_STATE = "hybrid"
BOXONLY_PROMPT_MODE = "Auto_box_fallback_dice_0p85_area_0p70_1p30"


DATASETS = [
    "ENID",
    "GLENDA",
    "GLENDA_clean",
]

SPLITS = [
    "val",
    "test",
]


METRICS = [
    "dice",
    "iou",
    "precision",
    "recall",
]


ALPHA = 0.05


# =============================================================================
# Path helpers
# =============================================================================

def get_posneg_metrics_path(dataset_key: str, split_key: str):
    return (
        RESULTS_ROOT
        / dataset_key
        / POSNEG_METHOD_FOLDER
        / POSNEG_TRAINING_STATE
        / POSNEG_PROMPT_MODE
        / split_key
        / "metrics_image_level.csv"
    )


def get_boxonly_metrics_path(dataset_key: str, split_key: str):
    return (
        RESULTS_ROOT
        / dataset_key
        / BOXONLY_METHOD_FOLDER
        / BOXONLY_TRAINING_STATE
        / BOXONLY_PROMPT_MODE
        / split_key
        / "metrics_image_level.csv"
    )


# =============================================================================
# Numeric helpers
# =============================================================================

def safe_float_series(series):
    return pd.to_numeric(series, errors="coerce")


def summarize_values(values):
    values = pd.to_numeric(pd.Series(values), errors="coerce").dropna()

    if len(values) == 0:
        return {
            "n": 0,
            "mean": np.nan,
            "std": np.nan,
            "median": np.nan,
            "q1": np.nan,
            "q3": np.nan,
            "min": np.nan,
            "max": np.nan,
        }

    return {
        "n": int(len(values)),
        "mean": float(values.mean()),
        "std": float(values.std(ddof=1)) if len(values) > 1 else 0.0,
        "median": float(values.median()),
        "q1": float(values.quantile(0.25)),
        "q3": float(values.quantile(0.75)),
        "min": float(values.min()),
        "max": float(values.max()),
    }


def format_mean_sd(mean_value, sd_value):
    if not np.isfinite(mean_value):
        return ""

    if not np.isfinite(sd_value):
        return f"{mean_value:.4f}"

    return f"{mean_value:.4f} ± {sd_value:.4f}"


def format_median_iqr(median_value, q1_value, q3_value):
    if not np.isfinite(median_value):
        return ""

    if not np.isfinite(q1_value) or not np.isfinite(q3_value):
        return f"{median_value:.4f}"

    return f"{median_value:.4f} [{q1_value:.4f}, {q3_value:.4f}]"


def get_case_id_series(df: pd.DataFrame):
    if "case_id" in df.columns:
        return df["case_id"].astype(str)

    if "image_name" in df.columns:
        return df["image_name"].astype(str).apply(lambda x: Path(x).stem)

    raise ValueError("Could not find case_id or image_name column.")


def ensure_case_id(df: pd.DataFrame):
    df = df.copy()
    df["__case_id__"] = get_case_id_series(df)
    return df


# =============================================================================
# Statistical tests
# =============================================================================

def run_wilcoxon_paired(method_a_values, method_b_values):
    """
    Paired Wilcoxon signed-rank test.

    Direction:
        difference = method_b - method_a

    Positive difference:
        method_b performs better.

    Negative difference:
        method_a performs better.
    """

    method_a_values = np.asarray(method_a_values, dtype=np.float64)
    method_b_values = np.asarray(method_b_values, dtype=np.float64)

    valid_mask = np.isfinite(method_a_values) & np.isfinite(method_b_values)

    method_a_values = method_a_values[valid_mask]
    method_b_values = method_b_values[valid_mask]

    differences = method_b_values - method_a_values
    nonzero_differences = differences[differences != 0]

    if len(method_a_values) == 0:
        return {
            "n_paired": 0,
            "n_nonzero_differences": 0,
            "wilcoxon_statistic": np.nan,
            "p_value": np.nan,
            "z_approx_abs": np.nan,
            "effect_size_r_abs": np.nan,
            "test_note": "No valid paired values.",
        }

    if len(nonzero_differences) == 0:
        return {
            "n_paired": int(len(method_a_values)),
            "n_nonzero_differences": 0,
            "wilcoxon_statistic": 0.0,
            "p_value": 1.0,
            "z_approx_abs": 0.0,
            "effect_size_r_abs": 0.0,
            "test_note": "All paired differences are zero.",
        }

    try:
        result = wilcoxon(
            method_b_values,
            method_a_values,
            alternative="two-sided",
            zero_method="wilcox",
            correction=False,
            mode="auto",
        )

        p_value = float(result.pvalue)
        statistic = float(result.statistic)

    except Exception as error:
        return {
            "n_paired": int(len(method_a_values)),
            "n_nonzero_differences": int(len(nonzero_differences)),
            "wilcoxon_statistic": np.nan,
            "p_value": np.nan,
            "z_approx_abs": np.nan,
            "effect_size_r_abs": np.nan,
            "test_note": f"Wilcoxon failed: {error}",
        }

    if p_value <= 0:
        z_abs = np.inf
    elif p_value >= 1:
        z_abs = 0.0
    else:
        z_abs = float(norm.isf(p_value / 2.0))

    if len(nonzero_differences) > 0 and np.isfinite(z_abs):
        effect_size_r_abs = float(z_abs / math.sqrt(len(nonzero_differences)))
    else:
        effect_size_r_abs = np.nan

    return {
        "n_paired": int(len(method_a_values)),
        "n_nonzero_differences": int(len(nonzero_differences)),
        "wilcoxon_statistic": statistic,
        "p_value": p_value,
        "z_approx_abs": z_abs,
        "effect_size_r_abs": effect_size_r_abs,
        "test_note": "OK",
    }


def holm_adjust_pvalues(p_values):
    p_values = np.asarray(p_values, dtype=np.float64)
    adjusted = np.full_like(p_values, np.nan, dtype=np.float64)

    valid_indices = np.where(np.isfinite(p_values))[0]

    if len(valid_indices) == 0:
        return adjusted

    valid_p = p_values[valid_indices]
    order = np.argsort(valid_p)

    sorted_indices = valid_indices[order]
    sorted_p = valid_p[order]

    m = len(sorted_p)
    adjusted_sorted = np.zeros(m, dtype=np.float64)

    running_max = 0.0

    for rank_index, p_value in enumerate(sorted_p):
        multiplier = m - rank_index
        adjusted_p = min(1.0, p_value * multiplier)
        running_max = max(running_max, adjusted_p)
        adjusted_sorted[rank_index] = running_max

    for sorted_position, original_index in enumerate(sorted_indices):
        adjusted[original_index] = adjusted_sorted[sorted_position]

    return adjusted


def decide_better(method_a_name, method_b_name, median_difference, mean_difference, adjusted_p):
    """
    Difference = method_b - method_a.
    """

    if not np.isfinite(adjusted_p):
        return "Unable to test"

    if adjusted_p >= ALPHA:
        if median_difference > 0:
            return f"{method_b_name} numerically higher, not significant"
        if median_difference < 0:
            return f"{method_a_name} numerically higher, not significant"
        return "No difference"

    if median_difference > 0:
        return f"{method_b_name} significantly better"

    if median_difference < 0:
        return f"{method_a_name} significantly better"

    if mean_difference > 0:
        return f"{method_b_name} significantly better by mean difference"

    if mean_difference < 0:
        return f"{method_a_name} significantly better by mean difference"

    return "Significant test but zero median/mean difference"


# =============================================================================
# Loading
# =============================================================================

def load_posneg_df(dataset_key: str, split_key: str):
    path = get_posneg_metrics_path(dataset_key, split_key)

    if not path.exists():
        raise FileNotFoundError(f"Missing PosNeg metrics file: {path}")

    df = pd.read_csv(path)
    df = ensure_case_id(df)

    return df, path


def load_boxonly_df(dataset_key: str, split_key: str):
    path = get_boxonly_metrics_path(dataset_key, split_key)

    if not path.exists():
        raise FileNotFoundError(f"Missing box-only hybrid metrics file: {path}")

    df = pd.read_csv(path)
    df = ensure_case_id(df)

    return df, path


# =============================================================================
# Optional summary columns
# =============================================================================

def extract_optional_summary_columns(df: pd.DataFrame, prefix: str):
    optional_columns = [
        "num_boxes",
        "accepted_surgisam2_components",
        "fallback_to_segformer_components",
        "acceptance_rate_components",
        "negative_point_fallback_offset_px",
        "dice_fallback_threshold",
        "acceptance_iou_threshold",
        "acceptance_dice_threshold",
        "min_area_ratio",
        "max_area_ratio",
        "dilation_radius_px",
    ]

    output = {}

    for column in optional_columns:
        key_mean = f"{prefix} {column} mean"
        key_median = f"{prefix} {column} median"

        if column not in df.columns:
            output[key_mean] = np.nan
            output[key_median] = np.nan
            continue

        values = pd.to_numeric(df[column], errors="coerce").dropna()

        if len(values) == 0:
            output[key_mean] = np.nan
            output[key_median] = np.nan
        else:
            output[key_mean] = float(values.mean())
            output[key_median] = float(values.median())

    return output


# =============================================================================
# Analysis helpers
# =============================================================================

def make_comparison_row(
    dataset_key,
    split_key,
    comparison_name,
    method_a_name,
    method_b_name,
    metric_name,
    method_a_values,
    method_b_values,
    posneg_path,
    boxonly_path,
    extra_summary,
):
    method_a_values = safe_float_series(method_a_values)
    method_b_values = safe_float_series(method_b_values)

    valid_mask = method_a_values.notna() & method_b_values.notna()

    method_a_values_valid = method_a_values[valid_mask]
    method_b_values_valid = method_b_values[valid_mask]
    differences = method_b_values_valid - method_a_values_valid

    a_summary = summarize_values(method_a_values_valid)
    b_summary = summarize_values(method_b_values_valid)
    diff_summary = summarize_values(differences)

    test_result = run_wilcoxon_paired(
        method_a_values=method_a_values_valid.values,
        method_b_values=method_b_values_valid.values,
    )

    improved_count = int((differences > 0).sum())
    worsened_count = int((differences < 0).sum())
    unchanged_count = int((differences == 0).sum())

    n_valid = int(len(differences))

    improvement_rate = improved_count / n_valid if n_valid > 0 else np.nan
    worsening_rate = worsened_count / n_valid if n_valid > 0 else np.nan
    unchanged_rate = unchanged_count / n_valid if n_valid > 0 else np.nan

    row = {
        "Comparison": comparison_name,
        "Dataset": dataset_key,
        "Split": split_key,
        "Metric": metric_name,

        "Method A": method_a_name,
        "Method B": method_b_name,
        "Difference definition": f"{method_b_name} - {method_a_name}",

        "n paired": test_result["n_paired"],
        "n nonzero differences": test_result["n_nonzero_differences"],

        f"{method_a_name} mean": a_summary["mean"],
        f"{method_a_name} SD": a_summary["std"],
        f"{method_a_name} mean ± SD": format_mean_sd(
            a_summary["mean"],
            a_summary["std"],
        ),
        f"{method_a_name} median": a_summary["median"],
        f"{method_a_name} Q1": a_summary["q1"],
        f"{method_a_name} Q3": a_summary["q3"],
        f"{method_a_name} median [Q1, Q3]": format_median_iqr(
            a_summary["median"],
            a_summary["q1"],
            a_summary["q3"],
        ),

        f"{method_b_name} mean": b_summary["mean"],
        f"{method_b_name} SD": b_summary["std"],
        f"{method_b_name} mean ± SD": format_mean_sd(
            b_summary["mean"],
            b_summary["std"],
        ),
        f"{method_b_name} median": b_summary["median"],
        f"{method_b_name} Q1": b_summary["q1"],
        f"{method_b_name} Q3": b_summary["q3"],
        f"{method_b_name} median [Q1, Q3]": format_median_iqr(
            b_summary["median"],
            b_summary["q1"],
            b_summary["q3"],
        ),

        "Mean difference MethodB-MethodA": diff_summary["mean"],
        "SD difference": diff_summary["std"],
        "Median difference MethodB-MethodA": diff_summary["median"],
        "Q1 difference": diff_summary["q1"],
        "Q3 difference": diff_summary["q3"],
        "Min difference": diff_summary["min"],
        "Max difference": diff_summary["max"],

        f"Images where {method_b_name} higher": improved_count,
        f"Images where {method_a_name} higher": worsened_count,
        "Unchanged images": unchanged_count,
        f"{method_b_name} improvement rate": improvement_rate,
        f"{method_a_name} higher rate": worsening_rate,
        "Unchanged rate": unchanged_rate,

        "Wilcoxon statistic": test_result["wilcoxon_statistic"],
        "Wilcoxon p": test_result["p_value"],
        "Holm-adjusted p": np.nan,
        "Significant after Holm": "",
        "Effect size r abs": test_result["effect_size_r_abs"],
        "z approx abs": test_result["z_approx_abs"],
        "Better method": "",
        "Test note": test_result["test_note"],

        "PosNeg metrics file": str(posneg_path),
        "BoxOnly metrics file": str(boxonly_path) if boxonly_path is not None else "",
    }

    row.update(extra_summary)

    return row, valid_mask, differences


def add_image_level_rows_from_same_df(
    output_rows,
    df,
    valid_mask,
    differences,
    dataset_key,
    split_key,
    comparison_name,
    method_a_name,
    method_b_name,
    metric_name,
    method_a_col,
    method_b_col,
):
    valid_indices = df[valid_mask].index.to_list()
    differences = list(differences)

    for index, difference_value in zip(valid_indices, differences):
        case_id = df.loc[index, "__case_id__"]
        image_name = df.loc[index, "image_name"] if "image_name" in df.columns else ""

        method_a_value = float(df.loc[index, method_a_col])
        method_b_value = float(df.loc[index, method_b_col])

        image_level_row = {
            "Comparison": comparison_name,
            "Dataset": dataset_key,
            "Split": split_key,
            "Metric": metric_name,
            "case_id": case_id,
            "image_name": image_name,
            "Method A": method_a_name,
            "Method B": method_b_name,
            "Method A value": method_a_value,
            "Method B value": method_b_value,
            "Difference MethodB-MethodA": difference_value,
            "Method B higher": "Yes" if difference_value > 0 else "No",
            "Method A higher": "Yes" if difference_value < 0 else "No",
            "Unchanged": "Yes" if difference_value == 0 else "No",
        }

        optional_columns = [
            "num_boxes",
            "accepted_surgisam2_components",
            "fallback_to_segformer_components",
            "acceptance_rate_components",
            "negative_point_fallback_offset_px",
            "dice_fallback_threshold",
        ]

        for optional_column in optional_columns:
            if optional_column in df.columns:
                image_level_row[f"PosNeg {optional_column}"] = df.loc[index, optional_column]

        output_rows.append(image_level_row)


def add_image_level_rows_from_merged_df(
    output_rows,
    merged_df,
    valid_mask,
    differences,
    dataset_key,
    split_key,
    comparison_name,
    method_a_name,
    method_b_name,
    metric_name,
    method_a_col,
    method_b_col,
):
    valid_indices = merged_df[valid_mask].index.to_list()
    differences = list(differences)

    for index, difference_value in zip(valid_indices, differences):
        case_id = merged_df.loc[index, "__case_id__"]
        image_name = ""

        if "image_name_posneg" in merged_df.columns:
            image_name = merged_df.loc[index, "image_name_posneg"]
        elif "image_name_boxonly" in merged_df.columns:
            image_name = merged_df.loc[index, "image_name_boxonly"]

        method_a_value = float(merged_df.loc[index, method_a_col])
        method_b_value = float(merged_df.loc[index, method_b_col])

        image_level_row = {
            "Comparison": comparison_name,
            "Dataset": dataset_key,
            "Split": split_key,
            "Metric": metric_name,
            "case_id": case_id,
            "image_name": image_name,
            "Method A": method_a_name,
            "Method B": method_b_name,
            "Method A value": method_a_value,
            "Method B value": method_b_value,
            "Difference MethodB-MethodA": difference_value,
            "Method B higher": "Yes" if difference_value > 0 else "No",
            "Method A higher": "Yes" if difference_value < 0 else "No",
            "Unchanged": "Yes" if difference_value == 0 else "No",
        }

        optional_columns = [
            "num_boxes_posneg",
            "accepted_surgisam2_components_posneg",
            "fallback_to_segformer_components_posneg",
            "acceptance_rate_components_posneg",
            "negative_point_fallback_offset_px_posneg",
            "dice_fallback_threshold_posneg",
            "num_boxes_boxonly",
            "accepted_surgisam2_components_boxonly",
            "fallback_to_segformer_components_boxonly",
            "acceptance_rate_components_boxonly",
        ]

        for optional_column in optional_columns:
            if optional_column in merged_df.columns:
                image_level_row[optional_column] = merged_df.loc[index, optional_column]

        output_rows.append(image_level_row)


# =============================================================================
# One dataset/split
# =============================================================================

def analyze_one_dataset_split(dataset_key: str, split_key: str):
    posneg_df, posneg_path = load_posneg_df(dataset_key, split_key)

    try:
        boxonly_df, boxonly_path = load_boxonly_df(dataset_key, split_key)
    except FileNotFoundError as error:
        print(f"WARNING: {error}")
        boxonly_df = None
        boxonly_path = None

    summary_rows = []
    image_level_rows = []

    posneg_optional_summary = extract_optional_summary_columns(
        posneg_df,
        prefix="PosNeg",
    )

    boxonly_optional_summary = {}
    if boxonly_df is not None:
        boxonly_optional_summary = extract_optional_summary_columns(
            boxonly_df,
            prefix="BoxOnly",
        )

    # -------------------------------------------------------------------------
    # Comparison 1:
    # SegFormer alone vs PosNeg hybrid
    #
    # Uses the SegFormer initial columns saved inside the PosNeg hybrid CSV.
    # -------------------------------------------------------------------------

    comparison_name = "SegFormer vs PosNeg hybrid"
    method_a_name = "SegFormer"
    method_b_name = "PosNeg hybrid"

    segformer_cols = {
        "dice": "segformer_initial_dice",
        "iou": "segformer_initial_iou",
        "precision": "segformer_initial_precision",
        "recall": "segformer_initial_recall",
    }

    for metric_name in METRICS:
        method_a_col = segformer_cols[metric_name]
        method_b_col = metric_name

        if method_a_col not in posneg_df.columns or method_b_col not in posneg_df.columns:
            print(
                f"WARNING: Missing columns in PosNeg file for "
                f"{dataset_key} | {split_key} | {metric_name}: "
                f"{method_a_col}, {method_b_col}"
            )
            continue

        extra_summary = {}
        extra_summary.update(posneg_optional_summary)

        row, valid_mask, differences = make_comparison_row(
            dataset_key=dataset_key,
            split_key=split_key,
            comparison_name=comparison_name,
            method_a_name=method_a_name,
            method_b_name=method_b_name,
            metric_name=metric_name,
            method_a_values=posneg_df[method_a_col],
            method_b_values=posneg_df[method_b_col],
            posneg_path=posneg_path,
            boxonly_path=boxonly_path,
            extra_summary=extra_summary,
        )

        summary_rows.append(row)

        add_image_level_rows_from_same_df(
            output_rows=image_level_rows,
            df=posneg_df,
            valid_mask=valid_mask,
            differences=differences,
            dataset_key=dataset_key,
            split_key=split_key,
            comparison_name=comparison_name,
            method_a_name=method_a_name,
            method_b_name=method_b_name,
            metric_name=metric_name,
            method_a_col=method_a_col,
            method_b_col=method_b_col,
        )

    # -------------------------------------------------------------------------
    # Comparison 2:
    # Box-only fallback hybrid vs PosNeg hybrid
    # -------------------------------------------------------------------------

    if boxonly_df is not None:
        comparison_name = "Box-only hybrid vs PosNeg hybrid"
        method_a_name = "Box-only hybrid"
        method_b_name = "PosNeg hybrid"

        merge_columns_posneg = [
            "__case_id__",
            "image_name",
            "dice",
            "iou",
            "precision",
            "recall",
            "num_boxes",
            "accepted_surgisam2_components",
            "fallback_to_segformer_components",
            "acceptance_rate_components",
            "negative_point_fallback_offset_px",
            "dice_fallback_threshold",
        ]

        merge_columns_boxonly = [
            "__case_id__",
            "image_name",
            "dice",
            "iou",
            "precision",
            "recall",
            "num_boxes",
            "accepted_surgisam2_components",
            "fallback_to_segformer_components",
            "acceptance_rate_components",
        ]

        merge_columns_posneg = [
            col for col in merge_columns_posneg
            if col in posneg_df.columns
        ]

        merge_columns_boxonly = [
            col for col in merge_columns_boxonly
            if col in boxonly_df.columns
        ]

        posneg_small = posneg_df[merge_columns_posneg].copy()
        boxonly_small = boxonly_df[merge_columns_boxonly].copy()

        merged_df = pd.merge(
            boxonly_small,
            posneg_small,
            on="__case_id__",
            how="inner",
            suffixes=("_boxonly", "_posneg"),
        )

        if len(merged_df) == 0:
            print(
                f"WARNING: No paired cases after merging box-only and PosNeg for "
                f"{dataset_key} | {split_key}"
            )
        else:
            print(
                f"Paired box-only vs PosNeg cases for {dataset_key} | {split_key}: "
                f"{len(merged_df)}"
            )

        for metric_name in METRICS:
            method_a_col = f"{metric_name}_boxonly"
            method_b_col = f"{metric_name}_posneg"

            if method_a_col not in merged_df.columns or method_b_col not in merged_df.columns:
                print(
                    f"WARNING: Missing merged columns for "
                    f"{dataset_key} | {split_key} | {metric_name}: "
                    f"{method_a_col}, {method_b_col}"
                )
                continue

            extra_summary = {}
            extra_summary.update(posneg_optional_summary)
            extra_summary.update(boxonly_optional_summary)

            row, valid_mask, differences = make_comparison_row(
                dataset_key=dataset_key,
                split_key=split_key,
                comparison_name=comparison_name,
                method_a_name=method_a_name,
                method_b_name=method_b_name,
                metric_name=metric_name,
                method_a_values=merged_df[method_a_col],
                method_b_values=merged_df[method_b_col],
                posneg_path=posneg_path,
                boxonly_path=boxonly_path,
                extra_summary=extra_summary,
            )

            summary_rows.append(row)

            add_image_level_rows_from_merged_df(
                output_rows=image_level_rows,
                merged_df=merged_df,
                valid_mask=valid_mask,
                differences=differences,
                dataset_key=dataset_key,
                split_key=split_key,
                comparison_name=comparison_name,
                method_a_name=method_a_name,
                method_b_name=method_b_name,
                metric_name=metric_name,
                method_a_col=method_a_col,
                method_b_col=method_b_col,
            )

    return summary_rows, image_level_rows


# =============================================================================
# Holm correction
# =============================================================================

def apply_holm_within_comparison_dataset_split(summary_df: pd.DataFrame):
    """
    Holm correction is applied within each:
        Comparison + Dataset + Split

    Example:
        ENID test, SegFormer vs PosNeg hybrid:
            Dice, IoU, precision, recall

    Those 4 p-values are Holm-corrected together.
    """

    summary_df = summary_df.copy()

    summary_df["Holm-adjusted p"] = np.nan
    summary_df["Significant after Holm"] = ""
    summary_df["Better method"] = ""

    group_cols = [
        "Comparison",
        "Dataset",
        "Split",
    ]

    for group_key, group_df in summary_df.groupby(group_cols):
        group_indices = group_df.index.to_list()
        p_values = group_df["Wilcoxon p"].values

        adjusted_values = holm_adjust_pvalues(p_values)

        for index, adjusted_p in zip(group_indices, adjusted_values):
            summary_df.loc[index, "Holm-adjusted p"] = adjusted_p

            significant = (
                np.isfinite(adjusted_p)
                and adjusted_p < ALPHA
            )

            summary_df.loc[index, "Significant after Holm"] = (
                "Yes" if significant else "No"
            )

            method_a_name = summary_df.loc[index, "Method A"]
            method_b_name = summary_df.loc[index, "Method B"]

            median_difference = summary_df.loc[
                index,
                "Median difference MethodB-MethodA",
            ]

            mean_difference = summary_df.loc[
                index,
                "Mean difference MethodB-MethodA",
            ]

            summary_df.loc[index, "Better method"] = decide_better(
                method_a_name=method_a_name,
                method_b_name=method_b_name,
                median_difference=median_difference,
                mean_difference=mean_difference,
                adjusted_p=adjusted_p,
            )

    return summary_df


# =============================================================================
# Interpretation/config sheets
# =============================================================================

def create_interpretation_sheet(summary_df: pd.DataFrame):
    rows = []

    rows.append(
        {
            "Item": "Purpose",
            "Description": (
                "This analysis tests whether the new SegFormer-guided SurgiSAM2 "
                "box+positive/negative point fallback hybrid improves over "
                "SegFormer alone and over the previous box-only fallback hybrid."
            ),
        }
    )

    rows.append(
        {
            "Item": "Comparison 1",
            "Description": (
                "SegFormer alone versus PosNeg hybrid. SegFormer values are read "
                "from the segformer_initial_* columns inside the PosNeg hybrid "
                "metrics_image_level.csv file."
            ),
        }
    )

    rows.append(
        {
            "Item": "Comparison 2",
            "Description": (
                "Box-only fallback hybrid versus PosNeg hybrid. The two CSV files "
                "are paired by case_id/image stem."
            ),
        }
    )

    rows.append(
        {
            "Item": "Pairing",
            "Description": (
                "All tests are paired by image. Each image contributes one value "
                "per method for each metric."
            ),
        }
    )

    rows.append(
        {
            "Item": "Difference",
            "Description": (
                "Difference = Method B - Method A. In both comparisons, Method B "
                "is the PosNeg hybrid. Positive values mean PosNeg hybrid improved "
                "the metric."
            ),
        }
    )

    rows.append(
        {
            "Item": "Statistical test",
            "Description": (
                "Two-sided Wilcoxon signed-rank test is used because results are "
                "paired by image and segmentation metrics are bounded and often "
                "non-normal."
            ),
        }
    )

    rows.append(
        {
            "Item": "Multiple testing correction",
            "Description": (
                "Holm correction is applied separately within each comparison, "
                "dataset, and split across Dice, IoU, precision, and recall."
            ),
        }
    )

    rows.append(
        {
            "Item": "Primary metric",
            "Description": (
                "Dice should be treated as the primary endpoint. IoU, precision, "
                "and recall are secondary metrics."
            ),
        }
    )

    rows.append(
        {
            "Item": "Decision rule",
            "Description": (
                "PosNeg hybrid is significantly better when the median difference "
                "is positive and Holm-adjusted p < 0.05. The comparator is "
                "significantly better when the median difference is negative and "
                "Holm-adjusted p < 0.05."
            ),
        }
    )

    if len(summary_df) > 0:
        dice_rows = summary_df[summary_df["Metric"] == "dice"].copy()

        for _, row in dice_rows.iterrows():
            rows.append(
                {
                    "Item": f"{row['Comparison']} | {row['Dataset']} {row['Split']} Dice conclusion",
                    "Description": (
                        f"{row['Better method']} | "
                        f"Method A={row['Method A']}, Method B={row['Method B']}, "
                        f"median difference={row['Median difference MethodB-MethodA']:.4f}, "
                        f"mean difference={row['Mean difference MethodB-MethodA']:.4f}, "
                        f"Holm-adjusted p={row['Holm-adjusted p']:.4g}."
                    ),
                }
            )

    return pd.DataFrame(rows)


def create_config_sheet():
    return pd.DataFrame(
        [
            {
                "Parameter": "RESULTS_ROOT",
                "Value": str(RESULTS_ROOT),
            },
            {
                "Parameter": "POSNEG_METHOD_FOLDER",
                "Value": POSNEG_METHOD_FOLDER,
            },
            {
                "Parameter": "POSNEG_TRAINING_STATE",
                "Value": POSNEG_TRAINING_STATE,
            },
            {
                "Parameter": "POSNEG_PROMPT_MODE",
                "Value": POSNEG_PROMPT_MODE,
            },
            {
                "Parameter": "BOXONLY_METHOD_FOLDER",
                "Value": BOXONLY_METHOD_FOLDER,
            },
            {
                "Parameter": "BOXONLY_TRAINING_STATE",
                "Value": BOXONLY_TRAINING_STATE,
            },
            {
                "Parameter": "BOXONLY_PROMPT_MODE",
                "Value": BOXONLY_PROMPT_MODE,
            },
            {
                "Parameter": "DATASETS",
                "Value": json.dumps(DATASETS),
            },
            {
                "Parameter": "SPLITS",
                "Value": json.dumps(SPLITS),
            },
            {
                "Parameter": "METRICS",
                "Value": json.dumps(METRICS),
            },
            {
                "Parameter": "ALPHA",
                "Value": ALPHA,
            },
            {
                "Parameter": "OUTPUT_XLSX",
                "Value": str(OUTPUT_XLSX),
            },
        ]
    )


# =============================================================================
# Excel writing
# =============================================================================

def autosize_worksheet_columns(writer, sheet_name: str, df: pd.DataFrame, max_width: int = 70):
    worksheet = writer.sheets[sheet_name]

    for column_index, column_name in enumerate(df.columns, start=1):
        column_values = (
            df[column_name]
            .astype(object)
            .where(pd.notna(df[column_name]), "")
            .astype(str)
        )

        max_len = max(
            [len(str(column_name))]
            + [len(str(value)) for value in column_values.head(200)]
        )

        width = min(max(max_len + 2, 10), max_width)
        column_letter = worksheet.cell(row=1, column=column_index).column_letter
        worksheet.column_dimensions[column_letter].width = width

def style_excel_sheet(writer, sheet_name: str, df: pd.DataFrame):
    worksheet = writer.sheets[sheet_name]
    worksheet.freeze_panes = "A2"

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    autosize_worksheet_columns(writer, sheet_name, df)


def write_xlsx(summary_df, image_level_df, interpretation_df, config_df):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        interpretation_df.to_excel(writer, sheet_name="interpretation", index=False)
        summary_df.to_excel(writer, sheet_name="summary_all", index=False)
        image_level_df.to_excel(writer, sheet_name="image_level_all", index=False)
        config_df.to_excel(writer, sheet_name="run_config", index=False)

        style_excel_sheet(writer, "interpretation", interpretation_df)
        style_excel_sheet(writer, "summary_all", summary_df)
        style_excel_sheet(writer, "image_level_all", image_level_df)
        style_excel_sheet(writer, "run_config", config_df)

        summary_sheet = writer.sheets["summary_all"]

        better_method_col = summary_df.columns.get_loc("Better method") + 1

        green_fill = PatternFill(
            start_color="C6EFCE",
            end_color="C6EFCE",
            fill_type="solid",
        )

        red_fill = PatternFill(
            start_color="FFC7CE",
            end_color="FFC7CE",
            fill_type="solid",
        )

        yellow_fill = PatternFill(
            start_color="FFEB9C",
            end_color="FFEB9C",
            fill_type="solid",
        )

        for row in range(2, summary_sheet.max_row + 1):
            better_method = summary_sheet.cell(
                row=row,
                column=better_method_col,
            ).value

            target_cell = summary_sheet.cell(
                row=row,
                column=better_method_col,
            )

            if isinstance(better_method, str) and "PosNeg hybrid significantly better" in better_method:
                target_cell.fill = green_fill

            elif isinstance(better_method, str) and "significantly better" in better_method:
                target_cell.fill = red_fill

            elif isinstance(better_method, str) and "not significant" in better_method:
                target_cell.fill = yellow_fill

    print(f"Saved XLSX: {OUTPUT_XLSX}")


# =============================================================================
# Main
# =============================================================================

def main():
    all_summary_rows = []
    all_image_level_rows = []

    for dataset_key in DATASETS:
        for split_key in SPLITS:
            print(f"Analyzing: {dataset_key} | {split_key}")

            try:
                summary_rows, image_level_rows = analyze_one_dataset_split(
                    dataset_key=dataset_key,
                    split_key=split_key,
                )

                all_summary_rows.extend(summary_rows)
                all_image_level_rows.extend(image_level_rows)

            except FileNotFoundError as error:
                print(f"WARNING: {error}")

    summary_df = pd.DataFrame(all_summary_rows)

    if len(summary_df) == 0:
        raise RuntimeError(
            "No summary rows were created. Check that PosNeg and box-only hybrid "
            "metrics_image_level.csv files exist."
        )

    summary_df = apply_holm_within_comparison_dataset_split(summary_df)

    image_level_df = pd.DataFrame(all_image_level_rows)

    interpretation_df = create_interpretation_sheet(summary_df)
    config_df = create_config_sheet()

    write_xlsx(
        summary_df=summary_df,
        image_level_df=image_level_df,
        interpretation_df=interpretation_df,
        config_df=config_df,
    )

    print("\nDONE.")
    print(f"Output: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()